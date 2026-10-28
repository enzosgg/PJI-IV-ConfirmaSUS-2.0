from django.shortcuts import render, HttpResponse
from django.http import JsonResponse
import openpyxl
from datetime import datetime
from apps.estatistica.logica import (
    grafico_total,
    grafico_agendada,
    grafico_compareceu,
    grafico_faltou,
    tabela,
    grafico_pizza,
    grafico_ofertada,
    grafico_absenteismo,
    grafico_perda_primaria,
    grafico_relacao,
    grafico_taxa_ocupacao,
)
import os
from django.conf import settings


def is_superuser(user):
    return user.is_authenticated and user.is_superuser


def estatistica_interface(request):
    dados = tabela()
    tabela_mensal = dados["mensal"]
    tabela_anual = dados["total_dados"]
    grafico1 = grafico_agendada(tabela_mensal)
    grafico2 = grafico_compareceu(tabela_mensal)
    grafico3 = grafico_faltou(tabela_mensal)
    grafico4 = grafico_absenteismo(tabela_mensal)
    grafico5 = grafico_ofertada(tabela_mensal)
    grafico6 = grafico_perda_primaria(tabela_mensal)
    grafico7 = grafico_relacao(tabela_mensal)
    grafico8 = grafico_taxa_ocupacao(tabela_mensal)
    grafico9 = grafico_total(tabela_anual)
    grafico10 = grafico_pizza(tabela_mensal)

    return render(
        request,
        "estatistica_interface.html",
        {
            "grafico_agendada": grafico1,
            "grafico_compareceu": grafico2,
            "grafico_faltou": grafico3,
            "grafico_absenteismo": grafico4,
            "grafico_ofertada": grafico5,
            "grafico_perda_primaria": grafico6,
            "grafico_relacao": grafico7,
            "grafico_taxa_ocupacao": grafico8,
            "grafico_total": grafico9,
            "grafico_pizza": grafico10,
            "tabela_interface": tabela_mensal,
            "tabela_total": tabela_anual,
        },
    )


def exportar_excel(request):
    dados = tabela()
    dados_mensais = dados["mensal"]
    workbook = openpyxl.load_workbook(
        filename=os.path.join(
            settings.BASE_DIR, "apps", "estatistica", "excel", "tabela.xlsx"
        )
    )
    sheet = workbook.active
    linha_inicial = 5
    coluna_inicial = 8

    for indice, dados_mes in enumerate(dados_mensais):
        linha_atual = linha_inicial + indice
        sheet.cell(row=linha_atual, column=coluna_inicial).value = dados_mes["ofertada"]
        sheet.cell(row=linha_atual, column=coluna_inicial + 1).value = dados_mes[
            "agendada"
        ]
        sheet.cell(row=linha_atual, column=coluna_inicial + 2).value = dados_mes[
            "compareceu"
        ]

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    nome_arquivo = f"relatorio_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    response["Content-Disposition"] = f"attachment; filename={nome_arquivo}"
    workbook.save(response)

    return response


def powerbi_dados_mensais(request):
    dados = tabela()
    dados_mensais = dados["mensal"]

    return JsonResponse(dados_mensais, safe=False)


def powerbi_dados_totais(request):
    dados = tabela()
    total_dados = dados["total_dados"]

    return JsonResponse(total_dados)
